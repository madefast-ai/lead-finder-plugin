#!/usr/bin/env python3
"""Lead Finder workbook: one .xlsx per campaign, one sheet per platform, generated overview sheets.

The platform sheets (LinkedIn, Companies, Reddit, Instagram) are the single source of truth. Every row has a
stable `key` (e.g. linkedin:pub:jane-doe). Lead Finder fills the source columns; the user (and Claude) own the
working columns: status, score, the interaction log, notes, drafts. Merging never overwrites a working column.

Lead Finder never reaches out cold. A lead goes new → to_engage → engaging (the user comments, reacts, connects)
→ warm (they engaged back) → talking (a conversation) → offer_sent → meeting → won or lost.

Summary, Next touches and Pipeline are rebuilt on every write. Don't edit them by hand.

Usage (python3; needs openpyxl: `pip install openpyxl`):
  workbook.py create   WORKBOOK
  workbook.py merge    WORKBOOK ROWS.json [--campaign NAME]   rows from get_results (JSON with "sheets") or a list
  workbook.py update   WORKBOOK KEY field=value [field=value ...]
  workbook.py interact WORKBOOK KEY --by you|them --channel linkedin|reddit|instagram|email --kind comment|like|connect|reply|message|… [--note TEXT] [--url POST_URL]
  workbook.py remove   WORKBOOK KEY_OR_URL                    delete a person everywhere in this workbook
  workbook.py import   WORKBOOK FILE.csv|FILE.xlsx [--sheet LinkedIn|Companies|Reddit|Instagram]
  workbook.py export   WORKBOOK OUT.csv [--sheet NAME] [--status STATUS[,STATUS]]
  workbook.py due      WORKBOOK [--days N]                   next touches due within N days (default 0: today and overdue)
  workbook.py summary  WORKBOOK
  workbook.py rebuild  WORKBOOK
Every command prints a JSON result.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import os
import re
import sys
import tempfile

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    print(json.dumps({"error": "openpyxl is missing. Run: pip install openpyxl"}))
    sys.exit(2)

TODAY = dt.date.today()
ENGAGE_DAYS = 3  # after you engage, look again in 3 days
REPLY_DAYS = 1  # when they engage back, answer by tomorrow
OFFER_DAYS = 4  # after an offer, follow up in 4 days

STATUSES = ["new", "to_engage", "engaging", "warm", "talking", "offer_sent", "meeting", "won", "lost", "not_interested", "do_not_contact"]
OPEN_STATUSES = {"engaging", "warm", "talking", "offer_sent", "meeting"}
CLOSED_STATUSES = {"won", "lost", "not_interested", "do_not_contact"}
NOT_ENGAGED = {"new", "to_engage", "engaging"}
# Statuses from older workbooks and CRM files.
LEGACY = {"to_contact": "to_engage", "contacted": "engaging", "replied": "talking"}
CONVERSATION_KINDS = {"message", "dm", "email", "call", "asked"}

# Columns the user and Claude own. Merges never overwrite them.
WORK = [
    "status", "score", "score_reason", "next_touch", "engage_channel", "touches", "engaged_on", "last_their_action", "last_their_action_on",
    "interactions", "notes", "draft_comment", "draft_message", "offer_sent_on", "offer_price", "offer_file", "campaign",
]
DATE_COLUMNS = {"next_touch", "engaged_on", "last_their_action_on", "offer_sent_on", "found_on"}

# Source and research columns Lead Finder fills (same order as the server's SHEET_COLUMNS).
SOURCE = {
    "LinkedIn": [
        "name", "headline", "role", "company", "company_url", "location", "profile_url", "open_to_work", "email", "email_quality",
        "signal", "signal_text", "signal_url", "signal_date", "source", "search", "found_on",
        "about", "experience", "education", "skills", "followers", "other_emails", "recent_posts", "recent_comments", "recent_reactions",
    ],
    "Companies": [
        "company", "linkedin_url", "industry", "size", "location", "website", "description",
        "signal", "signal_text", "signal_url", "signal_date", "source", "search", "found_on",
        "company_description", "company_website", "company_size", "company_industry", "company_hq", "company_founded", "company_specialities",
        "company_funding", "company_recent_posts",
    ],
    "Reddit": [
        "username", "profile_url", "subreddit", "type", "title", "text", "post_url", "posted_at", "upvotes", "comments",
        "signal", "source", "search", "found_on", "recent_activity", "karma", "about",
    ],
    "Instagram": [
        "username", "full_name", "profile_url", "caption", "hashtags", "post_url", "posted_at", "likes", "comments",
        "signal", "source", "search", "found_on", "bio", "followers", "website", "business_category", "business_account", "posts_count", "recent_posts",
    ],
}
SHEETS = list(SOURCE)

# What the user sees first on each sheet: who, then the working columns, then everything else.
FRONT = {
    "LinkedIn": ["key", "name", "role", "company"],
    "Companies": ["key", "company", "industry", "size"],
    "Reddit": ["key", "username", "subreddit", "title"],
    "Instagram": ["key", "username", "full_name", "followers"],
}
FRONT_WORK = ["status", "score", "score_reason", "next_touch", "engage_channel", "touches", "last_their_action", "notes"]
BACK_WORK = [c for c in WORK if c not in FRONT_WORK]
URL_COLUMNS = {"profile_url", "company_url", "signal_url", "linkedin_url", "website", "post_url", "company_website"}
GENERATED = ["Summary", "Next touches", "Pipeline"]
OLD_GENERATED = ["Follow-ups"]

# Headers people use in their own files → our columns (lower-case, spaces and punctuation removed).
ALIASES = {
    "fullname": "name", "firstname": "_first", "lastname": "_last", "title": "role", "jobtitle": "role", "position": "role",
    "companyname": "company", "organization": "company", "linkedin": "profile_url", "linkedinurl": "profile_url",
    "linkedinprofile": "profile_url", "profile": "profile_url", "url": "profile_url", "emailaddress": "email", "workemail": "email",
    "city": "location", "country": "location", "stage": "status", "leadstatus": "status", "comment": "notes", "comments_": "notes",
    "user": "username", "handle": "username", "instagram": "profile_url", "reddit": "profile_url", "companylinkedin": "company_url",
    "companyurl": "company_url", "domain": "website", "companywebsite": "website", "followup": "next_touch", "nextfollowup": "next_touch", "nexttouch": "next_touch",
}

HEADER_FILL = PatternFill("solid", fgColor="1D6B4F")
WORK_FILL = PatternFill("solid", fgColor="EAF3EE")


def out(obj) -> None:
    print(json.dumps(obj, ensure_ascii=False, indent=1, default=str))


def norm(h: str) -> str:
    return re.sub(r"[^a-z0-9_]", "", str(h or "").strip().lower().replace(" ", "_"))


def columns_for(sheet: str, extra=()) -> list[str]:
    front = FRONT[sheet]
    rest = [c for c in SOURCE[sheet] if c not in front]
    cols = front + FRONT_WORK + rest + BACK_WORK
    for e in extra:
        if e and e not in cols:
            cols.append(e)
    return cols


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    try:
        return dt.date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


# ---------- workbook I/O ----------


def load(path: str):
    if not os.path.exists(path):
        return create(path, save=False)
    wb = load_workbook(path)
    for s in SHEETS:
        if s not in wb.sheetnames:
            ws = wb.create_sheet(s)
            ws.append(columns_for(s))
    return wb


def create(path: str, save: bool = True):
    wb = Workbook()
    wb.remove(wb.active)
    for g in GENERATED:
        wb.create_sheet(g)
    for s in SHEETS:
        ws = wb.create_sheet(s)
        ws.append(columns_for(s))
    if save:
        write(wb, path)
    return wb


def read_rows(ws) -> tuple[list[str], list[dict]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    header = [str(h) if h is not None else "" for h in rows[0]]
    data = []
    for r in rows[1:]:
        if r is None or all(v is None or v == "" for v in r):
            continue
        data.append({header[i]: r[i] for i in range(min(len(header), len(r))) if header[i]})
    return header, data


def write_sheet(wb, sheet: str, rows: list[dict]) -> None:
    extra = []
    for r in rows:
        for k in r:
            if k not in extra:
                extra.append(k)
    cols = columns_for(sheet, extra)
    # Drop empty source columns so sheets stay readable; keep every working column.
    used = {k for r in rows for k, v in r.items() if v not in (None, "")}
    cols = [c for c in cols if c in used or c in WORK or c in FRONT[sheet]]
    idx = wb.sheetnames.index(sheet)
    wb.remove(wb[sheet])
    ws = wb.create_sheet(sheet, idx)
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c) for c in cols])
    style(ws, cols, with_status=True)


def style(ws, cols: list[str], with_status: bool = False) -> None:
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        letter = get_column_letter(i)
        long = c in {"about", "experience", "signal_text", "text", "caption", "recent_posts", "recent_comments", "recent_reactions", "notes",
                     "interactions", "draft_comment", "draft_message", "description", "company_description", "bio", "recent_activity", "score_reason"}
        ws.column_dimensions[letter].width = 48 if long else 22 if c in URL_COLUMNS else 14 if c in {"status", "score", "engage_channel", "touches"} else 18
        if c == "key":
            ws.column_dimensions[letter].hidden = True
        if c in WORK:
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=i).fill = WORK_FILL
        if c in URL_COLUMNS:
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=i)
                if isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = Font(color="1D6B4F", underline="single")
        if c in DATE_COLUMNS:
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=i).number_format = "yyyy-mm-dd"
    ws.freeze_panes = "C2" if ws.max_column > 2 else "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    if with_status and "status" in cols:
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
        col = get_column_letter(cols.index("status") + 1)
        dv.add(f"{col}2:{col}{max(ws.max_row, 2) + 500}")
        ws.add_data_validation(dv)


def write(wb, path: str) -> None:
    rebuild(wb)
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(os.path.abspath(path)) or ".")
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError:
        os.unlink(tmp)
        out({"error": f"Can't write {path}: close it in Excel/Numbers first, then try again."})
        sys.exit(3)


def all_rows(wb) -> dict[str, list[dict]]:
    return {s: read_rows(wb[s])[1] for s in SHEETS if s in wb.sheetnames}


# ---------- generated sheets ----------


def label(sheet: str, r: dict) -> str:
    return str(r.get("name") or r.get("username") or r.get("company") or r.get("full_name") or r.get("key") or "")


def rebuild(wb) -> None:
    data = all_rows(wb)
    for g in GENERATED + OLD_GENERATED:
        if g in wb.sheetnames:
            wb.remove(wb[g])

    # Summary
    ws = wb.create_sheet("Summary", 0)
    ws.append(["Lead Finder workbook"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Updated {dt.datetime.now().strftime('%Y-%m-%d %H:%M')}. Platform sheets are the source of truth; Summary, Next touches and Pipeline are rebuilt automatically."])
    ws.append(["Never message someone who hasn't interacted with you: engage in public first (comments, reactions). Delete anyone who asks."])
    ws.append([])
    ws.append(["Sheet"] + STATUSES + ["total"])
    for s in SHEETS:
        counts = [sum(1 for r in data.get(s, []) if (r.get("status") or "new") == st) for st in STATUSES]
        ws.append([s] + counts + [len(data.get(s, []))])
    due = touches_due(data, 0)
    ws.append([])
    ws.append(["Next touches due today or overdue", len(due)])
    for row in (5,):
        for c in ws[row]:
            c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 34

    # Next touches: due and planned, across platforms
    fu = wb.create_sheet("Next touches", 1)
    cols = ["due", "overdue_days", "sheet", "who", "status", "engage_channel", "last_their_action", "touches", "profile_url", "notes", "key"]
    fu.append(cols)
    for f in touches_due(data, 3650):
        fu.append([f.get(c) for c in cols])
    style(fu, cols)

    # Pipeline: everyone past "new", by status
    pl = wb.create_sheet("Pipeline", 2)
    cols = ["status", "sheet", "who", "score", "next_touch", "engaged_on", "engage_channel", "last_their_action", "profile_url", "key"]
    pl.append(cols)
    order = {s: i for i, s in enumerate(STATUSES)}
    rows = []
    for s, rs in data.items():
        for r in rs:
            st = r.get("status") or "new"
            if st in ("new",):
                continue
            rows.append([st, s, label(s, r), r.get("score"), to_date(r.get("next_touch")), to_date(r.get("engaged_on")), r.get("engage_channel"),
                         r.get("last_their_action"), r.get("profile_url") or r.get("linkedin_url"), r.get("key")])
    for row in sorted(rows, key=lambda x: (order.get(x[0], 99), x[4] or dt.date.max)):
        pl.append(row)
    style(pl, cols)


def touches_due(data: dict[str, list[dict]], days: int) -> list[dict]:
    limit = TODAY + dt.timedelta(days=days)
    out_rows = []
    for s, rs in data.items():
        for r in rs:
            st = r.get("status") or "new"
            d = to_date(r.get("next_touch"))
            if st in CLOSED_STATUSES or not d or d > limit:
                continue
            out_rows.append({
                "due": d, "overdue_days": max(0, (TODAY - d).days), "sheet": s, "who": label(s, r), "status": st, "engage_channel": r.get("engage_channel"),
                "last_their_action": r.get("last_their_action"), "touches": r.get("touches"), "profile_url": r.get("profile_url") or r.get("linkedin_url"),
                "notes": r.get("notes"), "key": r.get("key"),
            })
    return sorted(out_rows, key=lambda x: x["due"])


# ---------- commands ----------


def sheet_of_key(key: str, fallback: str | None = None) -> str:
    if key.startswith("linkedin:company"):
        return "Companies"
    if key.startswith("linkedin:"):
        return "LinkedIn"
    if key.startswith("reddit:"):
        return "Reddit"
    if key.startswith("instagram:"):
        return "Instagram"
    return fallback or "LinkedIn"


def incoming_rows(payload) -> list[tuple[str, dict]]:
    """get_results output ({"sheets": {...}}), a {"LinkedIn": [...]} map, or a plain list of rows with keys."""
    if isinstance(payload, dict) and "sheets" in payload:
        payload = payload["sheets"]
    rows = []
    if isinstance(payload, dict):
        for s, rs in payload.items():
            if s in SHEETS:
                rows += [(s, r) for r in rs]
    elif isinstance(payload, list):
        rows = [(sheet_of_key(str(r.get("key", ""))), r) for r in payload]
    return rows


def clean_incoming(r: dict) -> dict:
    r = {k: v for k, v in r.items() if k not in ("id", "state") and v not in (None, "")}
    if "score" in r and "upvotes" not in r and str(r.get("key", "")).startswith("reddit:"):
        r["upvotes"] = r.pop("score")  # Reddit's upvote score isn't the user's fit score
    return r


def merge_rows(wb, incoming: list[tuple[str, dict]], campaign: str | None, source_label: str | None = None) -> dict:
    data = all_rows(wb)
    added = updated = 0
    for s, raw in incoming:
        r = clean_incoming(raw)
        key = str(r.get("key") or "").strip()
        if not key:
            key = derive_key(s, r)
            if not key:
                continue
            r["key"] = key
        rows = data.setdefault(s, [])
        existing = next((x for x in rows if str(x.get("key")) == key), None)
        if existing is None:
            if r.get("profile_url"):
                existing = next((x for x in rows if x.get("profile_url") and str(x.get("profile_url")).rstrip("/").lower() == str(r["profile_url"]).rstrip("/").lower()), None)
        if existing is None:
            new = dict(r)
            new.setdefault("status", "new")
            new.setdefault("found_on", TODAY.isoformat())
            if campaign:
                new.setdefault("campaign", campaign)
            if source_label:
                new.setdefault("source", source_label)
            rows.append(new)
            added += 1
        else:
            changed = False
            for k, v in r.items():
                if k in WORK and existing.get(k) not in (None, ""):
                    continue  # never overwrite the user's working columns
                if k == "key":
                    continue
                if existing.get(k) != v:
                    existing[k] = v
                    changed = True
            updated += int(changed)
    for s, rows in data.items():
        write_sheet(wb, s, rows)
    return {"added": added, "updated": updated}


def derive_key(sheet: str, r: dict) -> str | None:
    url = str(r.get("profile_url") or r.get("linkedin_url") or "")
    m = re.search(r"linkedin\.com/in/([^/?#]+)", url, re.I)
    if m:
        slug = m.group(1)
        return f"linkedin:id:{slug}" if re.match(r"^AC[a-zA-Z0-9]AA", slug) else f"linkedin:pub:{slug.lower()}"
    m = re.search(r"linkedin\.com/company/([^/?#]+)", url, re.I)
    if m:
        return f"linkedin:company:{m.group(1).lower()}"
    m = re.search(r"reddit\.com/(?:u|user)/([^/?#]+)", url, re.I)
    if m or (sheet == "Reddit" and r.get("username")):
        return f"reddit:u:{(m.group(1) if m else str(r['username'])).lower()}"
    m = re.search(r"instagram\.com/([^/?#]+)", url, re.I)
    if m or (sheet == "Instagram" and r.get("username")):
        return f"instagram:u:{(m.group(1) if m else str(r['username']).lstrip('@')).lower()}"
    if sheet == "Companies" and r.get("company"):
        return f"linkedin:company-name:{str(r['company']).strip().lower()}"
    if r.get("email"):
        return f"email:{str(r['email']).strip().lower()}"
    return None


def find_row(data: dict[str, list[dict]], key: str):
    for s_, rows in data.items():
        for r in rows:
            if str(r.get("key")) == key or (r.get("profile_url") and str(r.get("profile_url")).rstrip("/") == key.rstrip("/")):
                return s_, rows, r
    return None, None, None


def apply_sets(r: dict, sets: dict) -> None:
    for k, v in sets.items():
        r[k] = to_date(v) if k in DATE_COLUMNS and v else (None if v == "" else v)


def in_days(n: int) -> str:
    return (TODAY + dt.timedelta(days=n)).isoformat()


def cmd_update(wb, key: str, pairs: list[str]) -> dict:
    data = all_rows(wb)
    sets = {}
    for p in pairs:
        if "=" not in p:
            raise SystemExit(f"Expected field=value, got {p}")
        k, v = p.split("=", 1)
        sets[k.strip()] = v
    s_, rows, r = find_row(data, key)
    if r is None:
        return {"error": f"No row with key or profile_url {key}"}
    st = sets.get("status")
    if st:
        st = sets["status"] = LEGACY.get(st, st)
        if st not in STATUSES:
            raise SystemExit(f"Unknown status {st}. Use one of: {', '.join(STATUSES)}")
        if st == "engaging" and not r.get("engaged_on"):
            sets.setdefault("engaged_on", TODAY.isoformat())
            sets.setdefault("next_touch", in_days(ENGAGE_DAYS))
        if st in ("warm", "talking"):
            sets.setdefault("next_touch", in_days(REPLY_DAYS))
        if st == "offer_sent":
            sets.setdefault("offer_sent_on", TODAY.isoformat())
            sets.setdefault("next_touch", in_days(OFFER_DAYS))
        if st in CLOSED_STATUSES:
            sets.setdefault("next_touch", "")
    apply_sets(r, sets)
    write_sheet(wb, s_, rows)
    return {"updated": key, "sheet": s_, "row": {k: r.get(k) for k in ["status", "next_touch", "engaged_on", "offer_sent_on", "notes"] if r.get(k) not in (None, "")}}


def cmd_interact(wb, key: str, by: str, channel: str, kind: str, note: str | None, url: str | None) -> dict:
    """Logs one interaction and moves the lead along: your touches make them engaging; theirs make them warm or talking."""
    if by not in ("you", "them"):
        raise SystemExit("--by must be you or them")
    data = all_rows(wb)
    s_, rows, r = find_row(data, key)
    if r is None:
        return {"error": f"No row with key or profile_url {key}"}
    st = LEGACY.get(r.get("status") or "new", r.get("status") or "new")
    line = f"{TODAY.isoformat()} {by}: {channel} {kind}" + (f" {url}" if url else "") + (f" · {note}" if note else "")
    log = [x for x in str(r.get("interactions") or "").split("\n") if x.strip()]
    sets: dict = {"interactions": "\n".join((log + [line])[-30:])}
    if by == "you":
        sets["touches"] = str(int(r.get("touches") or 0) + 1)
        sets["engage_channel"] = channel
        if not r.get("engaged_on"):
            sets["engaged_on"] = TODAY.isoformat()
        if st in ("new", "to_engage"):
            sets["status"] = "engaging"
        if st in NOT_ENGAGED:
            sets["next_touch"] = in_days(ENGAGE_DAYS)
    else:
        sets["last_their_action"] = f"{kind} ({channel})"
        sets["last_their_action_on"] = TODAY.isoformat()
        if kind in CONVERSATION_KINDS and st in NOT_ENGAGED | {"warm"}:
            sets["status"] = "talking"
        elif st in NOT_ENGAGED:
            sets["status"] = "warm"
        if st not in CLOSED_STATUSES:
            sets["next_touch"] = in_days(REPLY_DAYS)
    apply_sets(r, sets)
    write_sheet(wb, s_, rows)
    return {"logged": line, "sheet": s_, "row": {k: r.get(k) for k in ["status", "touches", "last_their_action", "next_touch"] if r.get(k) not in (None, "")}}


def cmd_remove(wb, needle: str) -> dict:
    data = all_rows(wb)
    removed = []
    n = needle.rstrip("/").lower()
    for s, rows in data.items():
        keep = []
        for r in rows:
            vals = {str(r.get(c) or "").rstrip("/").lower() for c in ("key", "profile_url", "email", "username")}
            if n in vals:
                removed.append({"sheet": s, "who": label(s, r)})
            else:
                keep.append(r)
        if len(keep) != len(rows):
            write_sheet(wb, s, keep)
    return {"removed": removed}


def read_any(path: str) -> list[dict]:
    if path.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            header, data = read_rows(ws)
            if header:
                rows += data
        return rows
    with open(path, newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return list(csv.DictReader(f, dialect=dialect))


def map_row(r: dict) -> dict:
    out_row: dict = {}
    extra: dict = {}
    for k, v in r.items():
        if v in (None, ""):
            continue
        n = norm(k)
        target = n if any(n in cols for cols in SOURCE.values()) or n in WORK or n == "key" else ALIASES.get(n.replace("_", ""))
        if target:
            if target in out_row:
                out_row[target] = f"{out_row[target]}, {v}" if target == "location" else out_row[target]
            else:
                out_row[target] = v
        else:
            extra[n or "column"] = v
    # "LinkedIn URL" columns usually hold people's profiles, not company pages.
    if "/in/" in str(out_row.get("linkedin_url", "")) and not out_row.get("profile_url"):
        out_row["profile_url"] = out_row.pop("linkedin_url")
    if "_first" in out_row or "_last" in out_row:
        out_row.setdefault("name", " ".join(str(out_row.get(x, "")) for x in ("_first", "_last")).strip())
    out_row.pop("_first", None)
    out_row.pop("_last", None)
    if isinstance(out_row.get("status"), str):
        s = out_row["status"].strip().lower().replace(" ", "_")
        s = LEGACY.get(s, s)
        out_row["status"] = s if s in STATUSES else "new"
    out_row.update({f"import_{k}": v for k, v in extra.items()})
    return out_row


def guess_sheet(r: dict) -> str:
    url = str(r.get("profile_url") or "")
    if "reddit.com" in url or r.get("subreddit"):
        return "Reddit"
    if "instagram.com" in url:
        return "Instagram"
    if "linkedin.com/company" in url or (r.get("company") and not r.get("name") and not r.get("username")):
        return "Companies"
    return "LinkedIn"


def cmd_import(wb, path: str, sheet: str | None) -> dict:
    rows = [map_row(r) for r in read_any(path)]
    incoming = [(sheet or guess_sheet(r), r) for r in rows]
    skipped = sum(1 for s, r in incoming if not (r.get("key") or derive_key(s, r)))
    res = merge_rows(wb, incoming, None, source_label=f"import: {os.path.basename(path)}")
    res["rows_in_file"] = len(rows)
    res["skipped_without_identifier"] = skipped
    return res


def cmd_export(wb, path: str, sheet: str | None, statuses: str | None) -> dict:
    data = all_rows(wb)
    wanted = set(statuses.split(",")) if statuses else None
    rows = []
    for s, rs in data.items():
        if sheet and s != sheet:
            continue
        for r in rs:
            if wanted and (r.get("status") or "new") not in wanted:
                continue
            rows.append({"sheet": s, **{k: v for k, v in r.items() if k != "key"}, "key": r.get("key")})
    cols: list[str] = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for r in rows:
            w.writerow([safe_cell(r.get(c)) for c in cols])
    return {"exported": len(rows), "file": path}


def safe_cell(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        s = v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(sep=" ", timespec="minutes")
    elif isinstance(v, dt.date):
        s = v.isoformat()
    else:
        s = str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@") else s


def cmd_summary(wb) -> dict:
    data = all_rows(wb)
    return {
        "sheets": {s: {st: n for st in STATUSES if (n := sum(1 for r in rs if (r.get("status") or "new") == st))} | {"total": len(rs)} for s, rs in data.items()},
        "touches_due": len(touches_due(data, 0)),
        "touches_next_7_days": len(touches_due(data, 7)),
    }


def main(argv: list[str]) -> None:
    if len(argv) < 3 or argv[1] in ("-h", "--help"):
        print(__doc__)
        sys.exit(1)
    cmd, path = argv[1], argv[2]
    opts = {argv[i]: argv[i + 1] for i in range(3, len(argv) - 1) if argv[i].startswith("--")}
    positional = [a for i, a in enumerate(argv[3:], start=3) if not a.startswith("--") and not argv[i - 1].startswith("--")]

    if cmd == "create":
        if os.path.exists(path):
            out({"exists": path})
            return
        create(path)
        out({"created": path, "sheets": GENERATED + SHEETS})
        return

    wb = load(path)
    if cmd == "merge":
        with open(positional[0], encoding="utf-8") as f:
            payload = json.load(f)
        res = merge_rows(wb, incoming_rows(payload), opts.get("--campaign"))
    elif cmd == "update":
        res = cmd_update(wb, positional[0], positional[1:])
    elif cmd == "interact":
        res = cmd_interact(wb, positional[0], opts.get("--by", ""), opts.get("--channel", "linkedin"), opts.get("--kind", "comment"), opts.get("--note"), opts.get("--url"))
    elif cmd == "remove":
        res = cmd_remove(wb, positional[0])
    elif cmd == "import":
        res = cmd_import(wb, positional[0], opts.get("--sheet"))
    elif cmd == "export":
        out(cmd_export(wb, positional[0], opts.get("--sheet"), opts.get("--status")))
        return
    elif cmd == "due":
        out({"due": touches_due(all_rows(wb), int(opts.get("--days", "0")))})
        return
    elif cmd == "summary":
        out(cmd_summary(wb))
        return
    elif cmd == "rebuild":
        res = {"rebuilt": True}
    else:
        print(__doc__)
        sys.exit(1)
    if "error" not in res:
        write(wb, path)
    out(res)


if __name__ == "__main__":
    main(sys.argv)
